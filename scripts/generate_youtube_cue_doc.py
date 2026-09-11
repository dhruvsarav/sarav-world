import json
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

JSON_PATH = Path(r'D:\Websites\SaravsWorld\public\apps\genzalphaslang\genz_alpha_lingo.json')
OUT_PATH_DOCS = Path(r'D:\sdrv\docs\Dhruv_GenZ_Alpha_Slang_YouTube_Practice_Script.docx')

# Custom actions mapped by ID
ACTIONS_MAP = {
    # Part 1: Lingo
    "lingo-01": "Pulls phone out fast, points camera down at screen, looks up with urgency: 'Yo, WYA? Drop your pin!'",
    "lingo-02": "Waves hand as if fanning flames because it's too hot, pursed lips: 'Sheesh, that fit is straight FIRE!'",
    "lingo-03": "Slow, respectful nod, hands held open in complete agreement: 'Honestly? Totally valid.'",
    "lingo-04": "Rhythmic head bob, shoulder bounce to an invisible beat: 'Wait... turn that up, that slaps!'",
    "lingo-05": "Casual shrug, unimpressed squint, slight mouth twist: 'Eh... it was just mid.'",
    "lingo-06": "Full-body grimace, pulls collar away from neck, squints in secondhand embarrassment: 'Bro, that is so cringe.'",
    "lingo-07": "Quick sharp nod, fist bump motion toward the camera: 'Bet. Say less.'",
    "lingo-08": "Pats chest over heart with flat palm, dead serious eye contact: 'No cap, 100% real.'",
    "lingo-09": "Taps imaginary baseball cap brim, shakes head smiling: 'Nah, stop the cap, bro!'",
    "lingo-10": "Slow-motion step back, hands shielding eyes from invisible radiant glow: 'His aura is over 9000!'",
    "lingo-11": "Freezes completely still, blank stare, stiff robotic walking cycle in place repeating one dialogue loop",
    "lingo-12": "Strikes an effortless high-fashion runway pose, hands on hips, head tilt: 'She ate and served!'",
    "lingo-13": "Rubs stomach, nods vigorously with sheer delight: 'Yo, this food is straight bussin!'",
    "lingo-14": "Makes 'W' sign with both hands (index and middle fingers joined), smiling wide: 'Huge W!'",
    "lingo-15": "Forms an 'L' with thumb and index finger, puts it gently to forehead with dramatic defeat: 'Took a massive L.'",
    "lingo-16": "Clutches chest dramatically, closes eyes in deep emotional appreciation: 'That song hits different.'",
    "lingo-17": "Points down toward the floor with both hands, shaking head in pity: 'Bro, please go outside and touch grass.'",
    "lingo-18": "Theatrical hand swoon over forehead like a dramatic actor fainting: 'Why are you being so extra?!'",
    "lingo-19": "Wide eyes, nervous chuckle, points finger spinning around ear: 'Bro has completely unhinged energy.'",
    "lingo-20": "Shivers with disgust, pulls shoulders in, makes a sour disgusted face: 'Ugh, immediate ick!'",
    "lingo-21": "Slow firm thumbs up, confident direct nod: 'Unapologetically based.'",
    "lingo-22": "Slaps cheeks lightly, eyes widen with razor focus, cracks knuckles: 'Time to lock in!'",
    "lingo-23": "Points up high with both hands, gasps in awe: 'This is literal PEAK fiction!'",
    "lingo-24": "Snaps fingers to rhythm, sways shoulders side-to-side with a cool smile: 'This song is an absolute bop!'",
    "lingo-25": "Slow-motion strut toward camera, adjusting sunglasses with invisible wind blowing",
    "lingo-26": "Shapes hands into binoculars over eyes, scanning room: 'Vibe check... you passed!'",
    "lingo-27": "Pinches waist in, strikes sharp hourglass silhouette pose with fierce eye contact: 'Looking completely snatched!'",
    "lingo-28": "Squints eyes suspiciously, creeps forward, points slowly at Dad: 'Hmm... that's mad sus.'",
    "lingo-29": "Taps side of forehead with index finger, smirks knowingly: 'Living in their head completely rent-free.'",
    "lingo-30": "Crisp military salute, then proud champion stance: 'Understood and delivered!'",
    "lingo-31": "Snaps finger down sharply, points index finger like a period mark: 'And that's that. Periodt!'",
    "lingo-32": "Pats Dad's shoulder sympathetically, shakes head slowly: 'I gotta be real with you chief... this ain't it.'",
    "lingo-33": "Slumps back in chair, exhales loudly, covers eyes with arm: 'Bro, we're totally cooked.'",
    "lingo-34": "Taps imaginary wristwatch, points directly at camera: 'You got clocked instantly!'",
    "lingo-35": "Smooth hair tuck behind ear, subtle wink at camera: 'Yeah, I just pulled.'",
    "lingo-36": "Hands on hips, sassy head roll, flips hair over shoulder with confidence: 'Certified baddie!'",
    "lingo-37": "Finger to lips 'shhh', looks left and right, winks: 'Keep it on the low, sneaky link.'",
    "lingo-38": "Mimics holding phone a foot away with reading glasses squint, typing with one index finger slowly like a boomer",
    "lingo-39": "Smooth hair sweep-back, silent confident smirk, points subtle finger gun without saying a word: 'Unspoken rizz.'",
    "lingo-40": "Hands grab hair, eyes bulge wildly, pacing frantically as if about to explode: 'Bro is about to crash out!'",
    "lingo-41": "Arms crossed tightly across chest, feet planted wide, resolute frown: 'We stand on business!'",
    "lingo-42": "Chef's kiss to the sky, wipes and dusts hands off cleanly: 'Ate down, left zero crumbs!'",
    "lingo-43": "Both hands palms-down making a 'hold on' calming gesture, leaning forward intently: 'Wait... let him cook!'",
    "lingo-44": "Shushes camera with index finger to lips, runs thumb firmly along jawline to highlight bone structure (Mewing)",
    "lingo-45": "Checks reflection in camera lens, tilts chin 45 degrees, fixes collar and hair meticulously: 'Looksmaxxing 101.'",
    "lingo-46": "Puffs out chest, chin lifted high, looks down nose with effortless elite billionaire smirk: 'Totally mogging them.'",
    "lingo-47": "Rapid exaggerated bowing, rubbing hands together in worship: 'Oh my god, you're the greatest! Stop glazing bro!'",
    "lingo-48": "Hand puppet mouth opening and closing rapidly like a duck quacking, looking bored at Dad: 'Bro is speaking fluent Yappanese!'",
    "lingo-49": "Keeps head perfectly still facing forward, darts eyes sharply sideways with one raised brow: 'Bombastic side eye... criminal offensive side eye.'",
    "lingo-50": "Ducks low, peers suspiciously around Dad's shoulder: 'Watch out, the opps are outside!'",
    "lingo-51": "Snaps fingers twice, sharp runway turn, striking an iconic pose: 'Slay all day!'",
    "lingo-52": "Smoothly adjusts jacket, flicks imaginary dust off shoulder, points down at sneakers: 'Check the drip!'",
    "lingo-53": "Wobbly hand gesture in the air, scratching head with puzzled look: 'It's... complicated. Just a situationship.'",
    "lingo-54": "Holds both hands out stopping Dad: 'Don't interfere, dad... it's a canon event, he has to learn.'",
    "lingo-55": "Eyes pop wide open, jaw drops in shock, hand flies to chest: 'GYAAAAT!'",
    "lingo-56": "Sneaky grin, reaches over smoothly and snatches a fry/snack straight from Dad's plate: 'Fanum tax, thank you!'",
    "lingo-57": "Playful bouncy head bob, goofy smirk, bouncy rhythm with finger guns: 'That is so skibidi!'",
    "lingo-58": "Taps temple with index finger, eyes swirling in mock dizziness: 'Too much scrolling, my brain is pure rot!'",
    "lingo-59": "Bewildered facepalm, slow head shake into camera: 'Bro... only in Ohio.'",
    "lingo-60": "Nods slowly, cold stoic expression, slight chin tuck, cool lone-wolf stare directly into the lens: 'Pure Sigma grindset.'",

    # Part 2: Idioms & Catchphrases
    "idioms-01": "Squints analytically, circles hand in air: 'It's giving... main character energy!'",
    "idioms-02": "Knocks gently on forehead, smirks: 'Not paying a single rupee of rent up there!'",
    "idioms-03": "Checks imaginary checklist with a pen, gives a crisp thumbs up: '10 out of 10, understood the assignment!'",
    "idioms-04": "Looks down at an imaginary plate, pats tummy: 'Cleaned the whole plate, left no crumbs!'",
    "idioms-05": "Pretends to squat and touch the ground, waving fresh air into face: 'Get some sunlight bro, touch grass!'",
    "idioms-06": "Holds hands up like a camera frame, clicks imaginary shutter: 'Caught you in 4K ultra HD!'",
    "idioms-07": "Leans in conspiratorially, taps side of nose: 'Show me the proof without saying a word!'",
    "idioms-08": "Quick guilty glance sideways, bites lip holding back laugh: 'Not me staying up till 3 AM scrolling TikTok!'",
    "idioms-09": "Puts on sunglasses, steps forward with open arms ready to party: 'We outside tonight!'",
    "idioms-10": "Slaps table or leg, crosses arms: 'I said what I said, standing on business!'",
    "idioms-11": "Pats air down to quiet everyone: 'Hold up, wait... let him cook!'",
    "idioms-12": "Wipes mouth with imaginary napkin, applauds smoothly: 'Completely ate that up!'",
    "idioms-13": "Unrolls long imaginary receipt down to floor, points: 'Look at line 3, I brought all the receipts!'",
    "idioms-14": "Cracks neck left and right, stares fiercely at camera: 'Time to lock in, focus mode ON!'",
    "idioms-15": "Sideways ice-skating hand motion across chest with smooth smile: 'Sliding right into the DMs.'",
    "idioms-16": "Flips hair, strikes dynamic superhero pose: 'The whole plot revolves around me, main character energy!'",
    "idioms-17": "Forms scanner with thumb and finger, scans Dad from head to toe: 'Beep... vibe check passed!'",
    "idioms-18": "Fakes sweet fluttery eyelashes, puts hands together: 'Oh, I'm just not like other kids! Pure pick-me.'",
    "idioms-19": "Drops jaw, pulls collar, staggers back dramatically like losing sanity: 'Bro is completely crashing out!'",
    "idioms-20": "Stops Dad by crossing arms into an 'X': 'No spoilers, it's a canon event, we cannot interfere.'",
    "idioms-21": "Taps head with a hopeful wide grin: 'Delusion is the only solution, trust the process! Delulu is the solulu!'",
    "idioms-22": "Presses imaginary side buttons on phone to take screenshot: 'Show the screenshot or it didn't happen!'",
    "idioms-23": "Dusts off fingertips, bow of satisfaction: 'Zero crumbs left behind!'",
    "idioms-24": "Knowing wink, small nod, finger pointed at audience: 'Real ones know. IYKYK.'",
    "idioms-25": "Throws imaginary paper airplane forward with full commitment: 'Just send it, full throttle!'",
    "idioms-26": "Quick smooth step slide to the side: 'I'll slide through in five minutes!'",
    "idioms-27": "Turns imaginary key in heart/pocket, puts key away: 'Locked down and secured!'",
    "idioms-28": "Dry sarcastic smile, slow patronizing clap: 'Oh wow... look at that, everybody's so creative!'",
    "idioms-29": "Flexes bicep, taps muscle: 'Putting in the reps every single day, reps on reps!'",
    "idioms-30": "Mimics defibrillator paddles: 'Clear! *BZZZT*', laughing: 'Reviving this dead group chat!'",
    "idioms-31": "Holds both hands out weighing scales up and down: 'Both sides are right, two things can be true at once!'",

    # Part 3: Shorthands & Abbreviations
    "abbr-01": "Deep double nod, flat hand to chest: 'FR bro, no joke, 100% real.'",
    "abbr-02": "Raises right hand swearing an oath: 'ONG, that actually happened!'",
    "abbr-03": "Covers mouth slightly, leans in whispered: 'NGL, that was actually pretty funny.'",
    "abbr-04": "Puts palms up in open honesty: 'TBH, I didn't even study for it.'",
    "abbr-05": "Snaps fingers, points enthusiastically at Dad: 'IKR?! That's literally what I just said!'",
    "abbr-06": "Taps ground twice with index finger: 'I need that right here, RN!'",
    "abbr-07": "Classic energetic shoulder shrug with palms out and tilted head: 'IDK man, you tell me!'",
    "abbr-08": "Pinch bridge of nose, slow disappointed head shake: 'SMH, what were you even thinking?'",
    "abbr-09": "Frames eyes with hands like movie director: 'POV: You forgot to do your homework and the teacher asks.'",
    "abbr-10": "Fast forward gesture with finger winding: 'Give me the TL;DR in five seconds flat.'",
    "abbr-11": "Stares into space with wistful dramatic expression: 'TFW the weekend is officially over.'",
    "abbr-12": "W-sign wave, turns head away dismissively: 'Whatever, w/e, do what you want.'",
    "abbr-13": "Cups hand over eyes looking around: 'WYA? We're all waiting outside!'",
    "abbr-14": "Pulls phone out, texts with one thumb, peeks up: 'Yo, WYD tonight? Any plans?'",
    "abbr-15": "Wrinkles forehead, tilts head in confusion: 'Wait, WYM by that? Explain.'",
    "abbr-16": "Hands out asking for clarification: 'Hold up, WDYM? That makes no sense.'",
    "abbr-17": "Points open palm toward Dad: 'That's what I think, WBU?'",
    "abbr-18": "Full 360 spin showing off sneakers and shirt: 'Check out today's OOTD!'",
    "abbr-19": "Pretends to comb hair and spray cologne: 'Get ready with me (GRWM) for the video shoot!'",
    "abbr-20": "Clutches stomach, bends over laughing silently: 'Bro, LMAO that is pure comedy!'",
    "abbr-21": "Bites fingernails anxiously checking phone: 'Serious FOMO watching everyone at the concert!'",
    "abbr-22": "Hands raised in self-defense: 'Hey, JMO, don't come at me, just saying!'",
    "abbr-23": "Touches chest, raises eyebrow: 'ISTG, I didn't touch your phone!'",
    "abbr-24": "Forms imaginary crown, places it gently on head with smug smile: 'Certified GOAT status.'",
    "abbr-25": "Silent knowing nod, tapping side of temple: 'IYKYK, no explanation needed for the real ones.'",
    "abbr-26": "Puts palm flat facing forward in 'STOP' sign: 'Phone on DND, I'm locking in.'",
    "abbr-27": "Points two fingers toward Dad: 'I'm doing great, HBU?'",
    "abbr-28": "Slaps hands together once, stares deadpan at Dad: 'Dad... BFFR, be for real right now!'",
    "abbr-29": "Steps out from behind imaginary screen: 'Wait till you actually meet them IRL, totally different.'",
    "abbr-30": "Punches fist in air: 'Double cheese pizza for dinner, FTW!'",
    "abbr-31": "Freezes with comical shocked expression: 'MFW Dad tells a pun in front of my friends.'",
    "abbr-32": "Covers ears with both hands, grimacing: 'Whoa whoa, TMI, that is way too much info!'",
    "abbr-33": "Walks two fingers across forearm like little legs walking away: 'Gotta step AFK for ten minutes!'",
    "abbr-34": "Gives a quick wave, dashes off-screen and pops right back: 'BRB in two seconds!'",
    "abbr-35": "Waves hand across chest as if erasing a chalkboard: 'Forget I asked, NVM.'",
    "abbr-36": "Offers respectful handshake to camera: 'GG bro, well played, respect.'",
    "abbr-37": "Snaps fingers, relaxed nod: 'OFC, anytime, you already know!'",
    "abbr-38": "Mimics running in place, glancing at watch: 'OMW right now, 5 mins away!'",
    "abbr-39": "Makes phone sign with hand (thumb to ear, pinky to mouth) and waves goodbye: 'TTYL, catch you later!'",
    "abbr-40": "Presses hands together in prayer gesture with grateful bow: 'TYSM, you're the absolute best!'",
    "abbr-41": "Gives finger guns with polite smile: 'TIA, appreciate you helping out!'",
    "abbr-42": "Taps chin thoughtfully: 'IMO, the sequel was ten times better.'",
    "abbr-43": "Touches heart politely, bows head slightly: 'IMHO, we should order from the new spot.'",
    "abbr-44": "Holds one hand flat balancing arguments: 'TBF, he did warn us three times.'",
    "abbr-45": "Brushes hands off like a job well done, winks: 'FTFY, fixed that for you, you're welcome!'",
    "abbr-46": "Head slowly drops forward onto table or hands: 'Dropped my phone in water... FML.'",
    "abbr-47": "Throws arms wide open, jumping slightly with pure excitement: 'YOLO! Let's do it right now!'",
    "abbr-48": "Points backwards with thumb: 'ICYMI, big news dropped on the channel yesterday!'",
    "abbr-49": "Points into camera lens with a mysterious grin: 'Shoutout to OOMF, they know who they are!'",
    "abbr-50": "Cups hands around mouth shouting, then points at viewer: 'SFS let's grow together!'",
    "abbr-51": "Gentle double pat in the air: 'DW bro, don't worry, I got this covered!'",
    "abbr-52": "Two thumbs up, wide gamer grin: 'GLHF everyone, may the best player win!'",
    "abbr-53": "Playful nudge, sticks tongue out with big laugh: 'JK JK, just kidding, don't get mad!'",
    "abbr-54": "Casual wave off, easy smile: 'NP at all, don't even mention it!'",
    "abbr-55": "Little finger wiggle: 'Not essential, just a sweet NTH feature.'",
    "abbr-56": "Makes heart shape with fingers over chest: 'They are the ultimate OTP, hands down!'",
    "abbr-57": "Pretends to lose balance laughing, grabs chair: 'ROFL that's too funny, I can't breathe!'",
    "abbr-58": "Puts hand on heart, raises eyebrow: 'STG, swear to god it wasn't me!'",
    "abbr-59": "Eyes light up, lightbulb gesture with finger clicking above head: 'TIL honey never spoils, wild!'",
    "abbr-60": "Pumps both fists in the air, sighs in relief: 'TGIF, finally the weekend is here!'",
    "abbr-61": "Mimics painting or hammering with focused expression: 'Don't judge yet, it's a WIP!'",
    "abbr-62": "Mimes texting on phone, points at Dad: 'HMU when you get free this evening!'",
    "abbr-63": "Sits back, crosses legs, opens arms welcomingly: 'Drop your questions in the comments, AMA!'",
    "abbr-64": "Tilts head with wide innocent child-like eyes: 'ELI5, explain like I'm five because that flew over my head!'",
    "abbr-65": "Quick nod, taps ear: 'LMK what you decide, I'm ready!'"
}

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=110, bottom=110, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, color="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def build_word_document():
    with open(JSON_PATH, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = Document()

    # Page Margins (0.65 inch for print efficiency)
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # Palette
    NAVY = RGBColor(15, 23, 42)       # #0F172A
    SLATE = RGBColor(71, 85, 105)     # #475569
    GOLD = RGBColor(180, 83, 9)       # #B45309
    GENZ_PURPLE = RGBColor(109, 40, 217)  # #6D28D9
    ALPHA_BLUE = RGBColor(2, 132, 199)    # #0284C7
    SHARED_TEAL = RGBColor(13, 148, 136)  # #0D9488
    ACTION_COLOR = RGBColor(180, 83, 9)   # Amber/Brown for acting cue

    # ==========================
    # HEADER / TITLE SECTION
    # ==========================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_super = title_p.add_run("🎬 YOUTUBE VIDEO RECORDING SCRIPT & PRACTICE CUE SHEET\n")
    run_super.font.name = "Segoe UI"
    run_super.font.size = Pt(9.5)
    run_super.font.bold = True
    run_super.font.color.rgb = GOLD

    run_title = title_p.add_run("THEN vs NOW: Millennial vs Gen Z & Alpha Slang Battle")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(2)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("A Complete 156-Term Rehearsal & Pronunciation Playbook for Sarav (Dad) & Dhruv (Son)")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SLATE

    # ==========================
    # CREATOR INSTRUCTIONS CALLOUT BOX
    # ==========================
    guide_table = doc.add_table(rows=1, cols=1)
    guide_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    guide_table.autofit = False
    guide_cell = guide_table.cell(0, 0)
    guide_cell.width = Inches(7.2)
    set_cell_background(guide_cell, "F1F5F9")
    set_cell_margins(guide_cell, top=140, bottom=140, left=180, right=180)

    # Border for callout box
    tcPr = guide_cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="B45309"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    gp = guide_cell.paragraphs[0]
    gp.paragraph_format.space_after = Pt(4)
    r = gp.add_run("📹 HOW TO SHOOT THIS YOUTUBE VIDEO (Quick Creator Guide):\n")
    r.font.name = "Segoe UI"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = NAVY

    tips = [
        ("1. Dynamic Two-Shot Format: ", "Dad (Sarav) delivers the Millennial line on the left with a puzzled or nostalgic face. Dhruv cuts in on the right with the modern slang, delivers the signature action/gesture, and drops the punchline!"),
        ("2. Batch Shooting: ", "Shoot in high-energy bursts of 10 to 15 terms per YouTube Short or Reel. Use the 3 categories (Lingo, Catchphrases, Abbreviations) as distinct video episodes."),
        ("3. Generation Badges: ", "[Gen Z] = Modern teen & young adult lingo | [Gen Alpha] = Internet-viral brainrot/gamer slang | [Shared] = Universal modern youth culture."),
        ("4. Practice Tip for Dhruv: ", "Exaggerate the physical action right as you say the slang word! The visual reaction sells the video.")
    ]
    for title_tip, desc in tips:
        tp = guide_cell.add_paragraph()
        tp.paragraph_format.space_before = Pt(2)
        tp.paragraph_format.space_after = Pt(2)
        r1 = tp.add_run(f"• {title_tip}")
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = tp.add_run(desc)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(9)
        r2.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==========================
    # SECTIONS & TABLES
    # ==========================
    section_icons = {
        "lingo": ("PART 1: VIRAL MODERN LINGO", "60 Terms — Core vocabulary, vibes, behavior, and dating dialect"),
        "idioms": ("PART 2: IDIOMS, CATCHPHRASES & COMEBACKS", "31 Phrases — Epic full-sentence comebacks, status checks, and internet idioms"),
        "abbreviations": ("PART 3: LIGHTNING CHAT SHORTHANDS", "65 Terms — Fast-finger DM abbreviations, gamer terms, and reactions")
    }

    for s_idx, sec in enumerate(data['sections']):
        sec_id = sec['id']
        sec_title, sec_desc = section_icons.get(sec_id, (sec['title'], ""))

        # Section Heading
        head_p = doc.add_paragraph()
        head_p.paragraph_format.space_before = Pt(16)
        head_p.paragraph_format.space_after = Pt(2)
        hrun = head_p.add_run(sec_title)
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(13)
        hrun.font.bold = True
        hrun.font.color.rgb = NAVY

        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_before = Pt(0)
        desc_p.paragraph_format.space_after = Pt(8)
        drun = desc_p.add_run(f"{sec_desc} ({len(sec['entries'])} total entries)")
        drun.font.name = "Segoe UI"
        drun.font.size = Pt(9.5)
        drun.font.italic = True
        drun.font.color.rgb = SLATE

        # 2-Column Table
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_row = table.rows[0]
        make_row_header(hdr_row)
        make_row_cant_split(hdr_row)

        c0 = hdr_row.cells[0]
        c1 = hdr_row.cells[1]
        c0.width = Inches(2.4)
        c1.width = Inches(4.8)
        set_cell_background(c0, "1E293B")  # Deep Navy
        set_cell_background(c1, "1E293B")
        set_cell_margins(c0, top=120, bottom=120, left=140, right=140)
        set_cell_margins(c1, top=120, bottom=120, left=140, right=140)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run("👴 MILLENNIAL ('THEN')\n[Dad's Setup Line]")
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run("⚡ GEN Z & ALPHA ('NOW') + 🎬 DHRUV'S ACTION\n[Slang Term, Physical Gesture & Dialogue Delivery]")
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(255, 255, 255)

        # Populate Entries
        for e_idx, entry in enumerate(sec['entries']):
            e_id = entry['id']
            row = table.add_row()
            make_row_cant_split(row)

            cell_left = row.cells[0]
            cell_right = row.cells[1]
            cell_left.width = Inches(2.4)
            cell_right.width = Inches(4.8)

            # Alternate row background
            bg_color = "F8FAFC" if (e_idx % 2 == 1) else "FFFFFF"
            set_cell_background(cell_left, bg_color)
            set_cell_background(cell_right, bg_color)
            set_cell_margins(cell_left, top=90, bottom=90, left=120, right=120)
            set_cell_margins(cell_right, top=90, bottom=90, left=120, right=120)
            set_cell_borders(cell_left, color="E2E8F0", sz="4")
            set_cell_borders(cell_right, color="E2E8F0", sz="4")

            # --- LEFT COLUMN (Dad / Millennial) ---
            pl = cell_left.paragraphs[0]
            pl.paragraph_format.space_before = Pt(1)
            pl.paragraph_format.space_after = Pt(2)

            num_run = pl.add_run(f"#{entry.get('order', e_idx+1)} ")
            num_run.font.name = "Segoe UI"
            num_run.font.size = Pt(8.5)
            num_run.font.bold = True
            num_run.font.color.rgb = GOLD

            then_run = pl.add_run(f'"{entry["then"]}"\n')
            then_run.font.name = "Segoe UI"
            then_run.font.size = Pt(10)
            then_run.font.bold = True
            then_run.font.color.rgb = NAVY

            ctx_run = pl.add_run(f"Context: {entry.get('context', '')}")
            ctx_run.font.name = "Segoe UI"
            ctx_run.font.size = Pt(8.5)
            ctx_run.font.color.rgb = SLATE

            # --- RIGHT COLUMN (Dhruv / GenZ Alpha) ---
            pr = cell_right.paragraphs[0]
            pr.paragraph_format.space_before = Pt(1)
            pr.paragraph_format.space_after = Pt(2)

            # Badge color
            gen = entry.get('generation', 'shared')
            gen_label = "Gen Z" if gen == "genz" else ("Gen Alpha" if gen == "alpha" else "Shared")
            gen_color = GENZ_PURPLE if gen == "genz" else (ALPHA_BLUE if gen == "alpha" else SHARED_TEAL)

            # Modern Slang Term
            now_run = pr.add_run(f"{entry['now']}  ")
            now_run.font.name = "Segoe UI"
            now_run.font.size = Pt(11)
            now_run.font.bold = True
            now_run.font.color.rgb = gen_color

            badge_run = pr.add_run(f"[{gen_label}]\n")
            badge_run.font.name = "Segoe UI"
            badge_run.font.size = Pt(8.5)
            badge_run.font.bold = True
            badge_run.font.color.rgb = gen_color

            # Action / Stage Direction
            action_desc = ACTIONS_MAP.get(e_id, f"Delivers line with confidence, smiles directly into camera, points finger for emphasis: '{entry.get('example', '')}'")
            act_p = cell_right.add_paragraph()
            act_p.paragraph_format.space_before = Pt(1)
            act_p.paragraph_format.space_after = Pt(2)

            act_icon = act_p.add_run("🎬 Preferred Action: ")
            act_icon.font.name = "Segoe UI"
            act_icon.font.size = Pt(8.5)
            act_icon.font.bold = True
            act_icon.font.color.rgb = ACTION_COLOR

            act_text = act_p.add_run(f"[{action_desc}]\n")
            act_text.font.name = "Segoe UI"
            act_text.font.size = Pt(9)
            act_text.font.italic = True
            act_text.font.color.rgb = RGBColor(51, 65, 85)

            # Example Dialogue line
            ex_p = cell_right.add_paragraph()
            ex_p.paragraph_format.space_before = Pt(0)
            ex_p.paragraph_format.space_after = Pt(2)

            ex_icon = ex_p.add_run("💬 Video Dialogue: ")
            ex_icon.font.name = "Segoe UI"
            ex_icon.font.size = Pt(8.5)
            ex_icon.font.bold = True
            ex_icon.font.color.rgb = NAVY

            ex_text = ex_p.add_run(f'"{entry.get("example", "")}"')
            ex_text.font.name = "Segoe UI"
            ex_text.font.size = Pt(9)
            ex_text.font.color.rgb = SLATE

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Save to private documents folder
    OUT_PATH_DOCS.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH_DOCS))

    print(f"Successfully generated Word Document for Dhruv:")
    print(f"-> {OUT_PATH_DOCS} ({OUT_PATH_DOCS.stat().st_size} bytes)")

if __name__ == "__main__":
    build_word_document()
